import textwrap
from pathlib import Path
from typing import Any
from xml.sax.saxutils import escape
from zipfile import ZIP_DEFLATED, ZipFile

from docx import Document
from PIL import Image, ImageDraw, ImageFont
from reportlab.lib.pagesizes import A4
from reportlab.pdfgen.canvas import Canvas

from app.schemas.export import ExportOptions
from app.services.mermaid import mindmap_to_mermaid


def render_qcm_pdf(qcm: dict[str, Any], path: Path, options: ExportOptions) -> None:
    pdf = _PDFWriter(path, options.title or qcm.get("title") or "QCM")
    pdf.heading(options.title or qcm.get("title") or "QCM")
    pdf.text(f"Difficulty: {qcm.get('difficulty', 'unknown')}")
    pdf.spacer()

    for question in qcm.get("questions", []):
        pdf.heading(f"{question.get('order_index', 0)}. {question.get('question_text', '')}", size=12)
        for answer in question.get("answers", []):
            marker = " [correct]" if options.include_answers and answer.get("is_correct") else ""
            pdf.text(f"{answer.get('label', '-')}. {answer.get('answer_text', '')}{marker}")
        if options.include_explanations and question.get("explanation"):
            pdf.text(f"Explanation: {question['explanation']}")
        if options.include_sources:
            _write_source_line(pdf.text, question)
        pdf.spacer()
    pdf.save()


def render_summary_pdf(summary: dict[str, Any], path: Path, options: ExportOptions) -> None:
    pdf = _PDFWriter(path, options.title or summary.get("title") or "Summary")
    pdf.heading(options.title or summary.get("title") or "Summary")
    pdf.text(summary.get("content") or "")
    if options.include_sources:
        pdf.spacer()
        pdf.text(f"Sources: {', '.join(summary.get('document_ids') or [])}")
    pdf.save()


def render_qcm_docx(qcm: dict[str, Any], path: Path, options: ExportOptions) -> None:
    doc = Document()
    doc.add_heading(options.title or qcm.get("title") or "QCM", level=1)
    doc.add_paragraph(f"Difficulty: {qcm.get('difficulty', 'unknown')}")
    for question in qcm.get("questions", []):
        doc.add_heading(f"{question.get('order_index', 0)}. {question.get('question_text', '')}", level=2)
        for answer in question.get("answers", []):
            marker = " [correct]" if options.include_answers and answer.get("is_correct") else ""
            doc.add_paragraph(f"{answer.get('label', '-')}. {answer.get('answer_text', '')}{marker}")
        if options.include_explanations and question.get("explanation"):
            doc.add_paragraph(f"Explanation: {question['explanation']}")
        if options.include_sources:
            _write_source_line(doc.add_paragraph, question)
    doc.save(path)


def render_qcm_xlsx(qcm: dict[str, Any], path: Path, options: ExportOptions) -> None:
    headers = [
        "etablissement",
        "formation",
        "promotion",
        "module",
        "element",
        "ouvrage",
        "chapitre_ouv",
        "question",
        "ordre_importance",
        "chapitre_element",
        "lettre",
        "choix",
        "Reponse (0,1)",
    ]
    context = _qcm_academic_context(qcm)
    document_titles = qcm.get("_document_titles_by_id") or {}
    rows: list[list[Any]] = [headers]

    for question in sorted(qcm.get("questions", []), key=lambda item: item.get("order_index") or 0):
        ouvrage = _question_document_title(question, document_titles)
        chapter_reference = _chapter_reference(question)
        chapter_excerpt = _chapter_excerpt(question)
        importance = _difficulty_to_importance(question.get("difficulty") or qcm.get("difficulty"))
        for answer in sorted(question.get("answers", []), key=lambda item: item.get("order_index") or 0):
            rows.append(
                [
                    context.get("establishment", ""),
                    context.get("formation", ""),
                    context.get("level", ""),
                    context.get("module", ""),
                    context.get("subject", ""),
                    ouvrage,
                    chapter_reference,
                    question.get("question_text") or "",
                    importance,
                    chapter_excerpt,
                    answer.get("label") or "",
                    answer.get("answer_text") or "",
                    1 if answer.get("is_correct") else 0,
                ],
            )

    _write_xlsx(path, rows)


def render_summary_docx(summary: dict[str, Any], path: Path, options: ExportOptions) -> None:
    doc = Document()
    doc.add_heading(options.title or summary.get("title") or "Summary", level=1)
    doc.add_paragraph(summary.get("content") or "")
    if options.include_sources:
        doc.add_paragraph(f"Sources: {', '.join(summary.get('document_ids') or [])}")
    doc.save(path)


def render_mermaid_code(mindmap: dict[str, Any], path: Path) -> str:
    code = mindmap_to_mermaid(mindmap)
    path.write_text(code, encoding="utf-8")
    return code


def render_mermaid_png(mindmap: dict[str, Any], path: Path) -> None:
    code = mindmap_to_mermaid(mindmap)
    title = mindmap.get("title") or "MindMap"
    tree = _parse_mermaid_mindmap(code, title)
    font = _load_font(18)
    small_font = _load_font(15)
    children = tree["children"]
    midpoint = (len(children) + 1) // 2
    left_nodes = list(reversed(children[:midpoint]))
    right_nodes = children[midpoint:]

    root_width = 290
    branch_width = 320
    child_width = 260
    root_gap = 420
    child_gap = 330
    margin = 180
    width = max(2200, int((margin + child_width + child_gap + branch_width + root_gap + root_width / 2) * 2))
    height = max(980, max(_branch_stack_height(left_nodes), _branch_stack_height(right_nodes), 360) + 260)

    image = Image.new("RGB", (width, height), color=(250, 248, 253))
    draw = ImageDraw.Draw(image)
    _draw_dot_mesh(draw, width, height)

    center_x = width // 2
    center_y = height // 2

    root_box = _draw_node_box(draw, center_x, center_y, tree["label"], font, primary=True, width=root_width)
    if left_nodes:
        _draw_branch_group(draw, left_nodes, "left", root_box, center_y, font, small_font)
    if right_nodes:
        _draw_branch_group(draw, right_nodes, "right", root_box, center_y, font, small_font)

    image.save(path)


def _draw_dot_mesh(draw: ImageDraw.ImageDraw, width: int, height: int) -> None:
    for x in range(32, width, 56):
        for y in range(30, height, 56):
            draw.ellipse((x, y, x + 2, y + 2), fill=(225, 213, 238))


def _draw_branch_group(
    draw: ImageDraw.ImageDraw,
    nodes: list[dict[str, Any]],
    side: str,
    root_box: tuple[int, int, int, int],
    center_y: int,
    font: ImageFont.ImageFont,
    small_font: ImageFont.ImageFont,
) -> None:
    if not nodes:
        return
    gap = 46
    block_heights = [_branch_block_height(node) for node in nodes]
    total_height = sum(block_heights) + gap * max(0, len(nodes) - 1)
    current_y = center_y - total_height // 2
    root_edge_x = root_box[0] if side == "left" else root_box[2]
    x = root_edge_x - 420 if side == "left" else root_edge_x + 420
    for node, block_height in zip(nodes, block_heights, strict=True):
        y = current_y + block_height // 2
        box = _draw_node_box(draw, x, y, node["label"], font, width=320)
        node_edge_x = box[2] if side == "left" else box[0]
        _draw_connector(draw, root_edge_x, center_y, node_edge_x, y)
        _draw_child_nodes(draw, node.get("children", [])[:5], side, box, y, small_font)
        current_y += block_height + gap


def _draw_child_nodes(
    draw: ImageDraw.ImageDraw,
    children: list[dict[str, Any]],
    side: str,
    parent_box: tuple[int, int, int, int],
    parent_y: int,
    font: ImageFont.ImageFont,
) -> None:
    if not children:
        return
    parent_edge_x = parent_box[0] if side == "left" else parent_box[2]
    x = parent_edge_x - 330 if side == "left" else parent_edge_x + 330
    step = 62
    start_y = parent_y - ((len(children) - 1) * step // 2)
    for index, child in enumerate(children):
        y = start_y + index * step
        box = _draw_node_box(draw, x, y, child["label"], font, width=260, small=True)
        child_edge_x = box[2] if side == "left" else box[0]
        _draw_connector(draw, parent_edge_x, parent_y, child_edge_x, y, muted=True)


def _branch_stack_height(nodes: list[dict[str, Any]]) -> int:
    if not nodes:
        return 0
    return sum(_branch_block_height(node) for node in nodes) + 46 * max(0, len(nodes) - 1)


def _branch_block_height(node: dict[str, Any]) -> int:
    children_count = min(len(node.get("children", [])), 5)
    return max(118, children_count * 62 + 36)


def _draw_connector(
    draw: ImageDraw.ImageDraw,
    x1: int,
    y1: int,
    x2: int,
    y2: int,
    muted: bool = False,
) -> None:
    color = (186, 146, 216) if muted else (148, 84, 190)
    mid_x = (x1 + x2) // 2
    draw.line((x1, y1, mid_x, y1, mid_x, y2, x2, y2), fill=color, width=2)


def _draw_node_box(
    draw: ImageDraw.ImageDraw,
    center_x: int,
    center_y: int,
    label: str,
    font: ImageFont.ImageFont,
    width: int = 280,
    primary: bool = False,
    small: bool = False,
) -> tuple[int, int, int, int]:
    lines = _wrap_lines(label, 24 if not small else 20)[:3]
    line_height = 16
    padding_x = 18 if not small else 14
    padding_y = 14 if not small else 10
    box_width = width
    box_height = padding_y * 2 + line_height * len(lines)
    x1 = center_x - box_width // 2
    y1 = center_y - box_height // 2
    x2 = center_x + box_width // 2
    y2 = center_y + box_height // 2
    fill = (134, 45, 190) if primary else ((245, 240, 249) if not small else (255, 255, 255))
    outline = (134, 45, 190) if primary else (215, 192, 232)
    text_fill = (255, 255, 255) if primary else (38, 29, 44)
    draw.rounded_rectangle((x1, y1, x2, y2), radius=22 if not small else 16, fill=fill, outline=outline, width=2)
    text_y = y1 + padding_y
    for line in lines:
        bbox = draw.textbbox((0, 0), line, font=font)
        text_x = x1 + padding_x if not primary else center_x - (bbox[2] - bbox[0]) // 2
        draw.text((text_x, text_y), line, fill=text_fill, font=font)
        text_y += line_height
    return (x1, y1, x2, y2)


def _load_font(size: int) -> ImageFont.ImageFont:
    for font_name in ("DejaVuSans.ttf", "Arial.ttf"):
        try:
            return ImageFont.truetype(font_name, size=size)
        except OSError:
            continue
    return ImageFont.load_default()


def _parse_mermaid_mindmap(code: str, fallback_title: str) -> dict[str, Any]:
    root = {"label": fallback_title or "MindMap", "children": []}
    stack: list[tuple[int, dict[str, Any]]] = [(-1, root)]
    for raw_line in code.splitlines():
        if not raw_line.strip() or raw_line.strip().lower() == "mindmap":
            continue
        indent = len(raw_line) - len(raw_line.lstrip(" "))
        label = _clean_mermaid_label(raw_line.strip())
        if not label:
            continue
        while len(stack) > 1 and indent <= stack[-1][0]:
            stack.pop()
        node = {"label": label, "children": []}
        stack[-1][1]["children"].append(node)
        stack.append((indent, node))
    if len(root["children"]) == 1 and root["children"][0]["children"]:
        return root["children"][0]
    return root if root["children"] else {"label": fallback_title or "MindMap", "children": []}


def _clean_mermaid_label(value: str) -> str:
    cleaned = value.strip()
    if cleaned.lower().startswith("root"):
        cleaned = cleaned[4:].strip()
    cleaned = cleaned.strip("()[]{}\" ")
    return cleaned[:90] or "Item"


def _count_visible_nodes(node: dict[str, Any]) -> int:
    return 1 + sum(_count_visible_nodes(child) for child in node.get("children", [])[:8])


def _qcm_academic_context(qcm: dict[str, Any]) -> dict[str, str]:
    settings = qcm.get("settings") or {}
    custom_rules = settings.get("custom_rules") or {}
    metadata = custom_rules.get("advanced_metadata") or {}
    if not isinstance(metadata, dict):
        return {}
    return {str(key): str(value) for key, value in metadata.items() if value is not None}


def _question_document_title(question: dict[str, Any], document_titles: dict[str, str]) -> str:
    document_id = question.get("source_document_id")
    if document_id and str(document_id) in document_titles:
        return document_titles[str(document_id)]
    if document_titles:
        return next(iter(document_titles.values()))
    return ""


def _chapter_reference(question: dict[str, Any]) -> str:
    if question.get("source_page"):
        return f"Page {question['source_page']}"
    return ""


def _chapter_excerpt(question: dict[str, Any]) -> str:
    return str(question.get("citation") or "").strip()[:300]


def _difficulty_to_importance(value: Any) -> int | str:
    normalized = str(value or "").strip().lower()
    mapping = {
        "easy": 1,
        "facile": 1,
        "medium": 2,
        "moyen": 2,
        "hard": 3,
        "difficile": 3,
        "expert": 4,
    }
    return mapping.get(normalized, str(value or ""))


def _write_xlsx(path: Path, rows: list[list[Any]]) -> None:
    with ZipFile(path, "w", ZIP_DEFLATED) as archive:
        archive.writestr("[Content_Types].xml", _content_types_xml())
        archive.writestr("_rels/.rels", _root_rels_xml())
        archive.writestr("xl/workbook.xml", _workbook_xml())
        archive.writestr("xl/_rels/workbook.xml.rels", _workbook_rels_xml())
        archive.writestr("xl/worksheets/sheet1.xml", _sheet_xml(rows))
        archive.writestr("docProps/core.xml", _core_props_xml())
        archive.writestr("docProps/app.xml", _app_props_xml())


def _sheet_xml(rows: list[list[Any]]) -> str:
    row_xml = []
    for row_index, row in enumerate(rows, start=1):
        cells = "".join(_xlsx_cell(row_index, column_index, value) for column_index, value in enumerate(row, start=1))
        row_xml.append(f'<row r="{row_index}">{cells}</row>')
    return (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main" '
        'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships">'
        '<cols>'
        '<col min="1" max="5" width="24" customWidth="1"/>'
        '<col min="6" max="7" width="28" customWidth="1"/>'
        '<col min="8" max="8" width="60" customWidth="1"/>'
        '<col min="9" max="13" width="20" customWidth="1"/>'
        "</cols>"
        f"<sheetData>{''.join(row_xml)}</sheetData>"
        "</worksheet>"
    )


def _xlsx_cell(row_index: int, column_index: int, value: Any) -> str:
    reference = f"{_column_letter(column_index)}{row_index}"
    if isinstance(value, (int, float)):
        return f'<c r="{reference}"><v>{value}</v></c>'
    text = escape(str(value or ""))
    return f'<c r="{reference}" t="inlineStr"><is><t xml:space="preserve">{text}</t></is></c>'


def _column_letter(index: int) -> str:
    letters = ""
    while index:
        index, remainder = divmod(index - 1, 26)
        letters = chr(65 + remainder) + letters
    return letters


def _content_types_xml() -> str:
    return (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
        '<Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>'
        '<Default Extension="xml" ContentType="application/xml"/>'
        '<Override PartName="/xl/workbook.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml"/>'
        '<Override PartName="/xl/worksheets/sheet1.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.worksheet+xml"/>'
        '<Override PartName="/docProps/core.xml" ContentType="application/vnd.openxmlformats-package.core-properties+xml"/>'
        '<Override PartName="/docProps/app.xml" ContentType="application/vnd.openxmlformats-officedocument.extended-properties+xml"/>'
        "</Types>"
    )


def _root_rels_xml() -> str:
    return (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
        '<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="xl/workbook.xml"/>'
        '<Relationship Id="rId2" Type="http://schemas.openxmlformats.org/package/2006/relationships/metadata/core-properties" Target="docProps/core.xml"/>'
        '<Relationship Id="rId3" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/extended-properties" Target="docProps/app.xml"/>'
        "</Relationships>"
    )


def _workbook_xml() -> str:
    return (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<workbook xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main" '
        'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships">'
        '<sheets><sheet name="QCM" sheetId="1" r:id="rId1"/></sheets>'
        "</workbook>"
    )


def _workbook_rels_xml() -> str:
    return (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
        '<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/worksheet" Target="worksheets/sheet1.xml"/>'
        "</Relationships>"
    )


def _core_props_xml() -> str:
    return (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<cp:coreProperties xmlns:cp="http://schemas.openxmlformats.org/package/2006/metadata/core-properties" '
        'xmlns:dc="http://purl.org/dc/elements/1.1/">'
        "<dc:creator>Iktibar</dc:creator>"
        "<cp:lastModifiedBy>Iktibar</cp:lastModifiedBy>"
        "</cp:coreProperties>"
    )


def _app_props_xml() -> str:
    return (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<Properties xmlns="http://schemas.openxmlformats.org/officeDocument/2006/extended-properties" '
        'xmlns:vt="http://schemas.openxmlformats.org/officeDocument/2006/docPropsVTypes">'
        "<Application>Iktibar</Application>"
        "</Properties>"
    )


class _PDFWriter:
    def __init__(self, path: Path, title: str):
        self.canvas = Canvas(str(path), pagesize=A4)
        self.width, self.height = A4
        self.y = self.height - 48
        self.canvas.setTitle(title)

    def heading(self, value: str, size: int = 15) -> None:
        self._ensure_space(44)
        self.canvas.setFont("Helvetica-Bold", size)
        for line in _wrap_lines(value, 88):
            self.canvas.drawString(48, self.y, line)
            self.y -= size + 5
        self.canvas.setFont("Helvetica", 10)

    def text(self, value: str) -> None:
        self._ensure_space(26)
        self.canvas.setFont("Helvetica", 10)
        for line in _wrap_lines(value, 100):
            self._ensure_space(18)
            self.canvas.drawString(48, self.y, line)
            self.y -= 14

    def spacer(self) -> None:
        self.y -= 10

    def save(self) -> None:
        self.canvas.save()

    def _ensure_space(self, needed: int) -> None:
        if self.y - needed < 48:
            self.canvas.showPage()
            self.y = self.height - 48


def _wrap_lines(value: str, width: int) -> list[str]:
    lines: list[str] = []
    for raw_line in str(value).splitlines() or [""]:
        lines.extend(textwrap.wrap(raw_line, width=width) or [""])
    return lines


def _write_source_line(writer, question: dict[str, Any]) -> None:
    parts = []
    if question.get("source_page"):
        parts.append(f"page {question['source_page']}")
    if question.get("citation"):
        parts.append(str(question["citation"]))
    if parts:
        writer("Source: " + " - ".join(parts))
